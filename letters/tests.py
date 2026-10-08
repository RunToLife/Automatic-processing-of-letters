"""Тесты преобразования PDF -> HTML/DOCX: таблицы и расположение текста должны совпадать с оригиналом."""
import io
import json
import shutil
from unittest import skipUnless
from unittest import mock

import pymupdf
from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document
from PIL import Image

from . import layout as L
from . import ocrprep
from .converter import ConversionError, convert_pdf
from .docxbuild import html_to_docx

HAS_TESSERACT = shutil.which('tesseract') is not None
HAS_OCRMYPDF = ocrprep.available() and ocrprep.diagnostics()['ghostscript']


def _text(page, x, y, s, size=12, bold=False, align='l'):
    font = 'tibo' if bold else 'tiro'
    w = pymupdf.get_text_length(s, fontname=font, fontsize=size)
    x = x if align == 'l' else (x - w / 2 if align == 'c' else x - w)
    page.insert_text((x, y), s, fontname=font, fontsize=size)


def _letter(grid=True, lineless=False, ruled=False):
    """Письмо A4: шапка по центру, реквизиты слева / адресат справа, таблица, подпись «должность — ФИО»."""
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    _text(p, 160, 70, 'MINISTRY OF FINANCE', 11, True, 'c')
    _text(p, 60, 150, 'No. 15-123 from 12.05.2025')
    for i, line in enumerate(['To the Director', 'Ivanov Ivan Ivanovich', 'Moscow, Lenin st. 1']):
        _text(p, 340, 150 + i * 15, line)
    _text(p, 297, 215, 'About the supplies', 13, True, 'c')
    _text(p, 95, 250, 'We inform you about the delivery terms.')
    if grid:
        x = [60, 95, 285, 345, 405, 535]
        top = 290
        p.draw_rect(pymupdf.Rect(x[0], top, x[-1], top + 76), width=.8)
        for yy in (top + 22, top + 40, top + 58):
            p.draw_line((x[0], yy), (x[-1], yy), width=.8)
        for xi in x[1:-1]:
            p.draw_line((xi, top), (xi, top + 76), width=.8)
        for xx, s in ((77, 'No'), (190, 'Name'), (315, 'Unit'), (375, 'Qty'), (470, 'Sum')):
            _text(p, xx, top + 15, s, 10, True, 'c')
        for r, (a, b, c, d, e) in enumerate([('1', 'Laptop', 'pcs', '5', '425000'), ('2', 'Monitor', 'pcs', '10', '210000'),
                                             ('3', 'Cable', 'pcs', '20', '8000')]):
            yy = top + 35 + r * 18
            _text(p, 77, yy, a, 10, align='c')
            _text(p, 100, yy, b, 10)
            _text(p, 315, yy, c, 10, align='c')
            _text(p, 398, yy, d, 10, align='r')
            _text(p, 528, yy, e, 10, align='r')
    if lineless:
        for r, row in enumerate([('Stage', 'Owner', 'Deadline', 'State'), ('Analysis', 'Ivanov', '10.06', 'done'),
                                 ('Design', 'Petrov', '25.06', 'active'), ('Build', 'Sidorov', '30.07', 'waiting')]):
            for xx, s in zip((60, 220, 340, 450), row):
                _text(p, xx, 400 + r * 18, s, 10, bold=(r == 0))
    if ruled:
        top = 500
        for yy in (top, top + 20, top + 38, top + 56, top + 74):
            p.draw_line((60, yy), (535, yy), width=.8)
        for r, row in enumerate([('Stage', 'Deadline', 'State'), ('Delivery', '15.05', 'done'), ('Setup', '30.05', 'active'),
                                 ('Accept', '10.06', 'waiting')]):
            for xx, s in zip((62, 300, 400), row):
                _text(p, xx, top + 14 + r * 18, s, 10, bold=(r == 0))
    _text(p, 60, 700, 'Deputy Minister', 12, True)
    _text(p, 535, 700, 'A.G. Siluanov', 12, True, 'r')
    return doc


def _scan_of(doc, dpi=300):
    """Растровый PDF без текстового слоя (как скан)."""
    out = pymupdf.open()
    for pg in doc:
        pix = pg.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY, alpha=False)
        np_ = out.new_page(width=pg.rect.width, height=pg.rect.height)
        np_.insert_image(np_.rect, stream=pix.tobytes('png'))
    return out


def _convert(doc):
    return convert_pdf(doc.tobytes())


class LayoutHtmlTests(SimpleTestCase):
    def test_grid_table_becomes_table(self):
        res = _convert(_letter())
        self.assertIn('class="gd-grid"', res.html)
        grid = res.html[res.html.index('class="gd-grid"'):]
        grid = grid[:grid.index('</table>')]
        self.assertEqual(grid.count('<tr'), 4)
        for word in ('Laptop', 'Monitor', 'Cable', '425000'):
            self.assertIn(word, res.html)

    def test_requisites_and_addressee_are_side_by_side(self):
        res = _convert(_letter(grid=False))
        # «реквизиты | адресат» — невидимая таблица из двух колонок, а не один абзац
        self.assertIn('class="gd-layout"', res.html)
        left, right = res.html.index('No. 15-123'), res.html.index('To the Director')
        row = res.html[res.html.rindex('<table', 0, left):]
        self.assertLess(row.index('No. 15-123'), row.index('To the Director'))
        self.assertGreater(right, left)

    def test_centered_title_and_signature_row(self):
        res = _convert(_letter(grid=False))
        self.assertRegex(res.html, r'text-align:center[^"]*"><b>About the supplies</b>')
        sig = res.html[res.html.index('Deputy Minister') - 400:]
        self.assertIn('text-align:right', sig)          # ФИО справа

    def test_lineless_table_is_real_table_with_columns(self):
        res = _convert(_letter(grid=False, lineless=True))
        block = res.html[res.html.index('Stage'):]
        self.assertEqual(block[:block.index('Build')].count('<td'), 4 * 3)   # 4 колонки × 3 строки до «Build»
        self.assertIn('gd-layout', res.html)

    def test_horizontal_lines_table_keeps_borders(self):
        res = _convert(_letter(grid=False, ruled=True))
        self.assertIn('gd-custom', res.html)
        self.assertIn('border-bottom:0.5pt solid #000', res.html)
        self.assertIn('border-left:none', res.html)

    def test_page_meta_follows_content(self):
        res = _convert(_letter())
        self.assertAlmostEqual(res.meta['w'], 595, delta=1)
        self.assertAlmostEqual(res.meta['ml'], 60, delta=2)


class DocxTests(SimpleTestCase):
    def _docx(self, **kw):
        res = _convert(_letter(**kw))
        return Document(io.BytesIO(html_to_docx(res.html, res.meta))), res

    def test_docx_has_tables_and_page_geometry(self):
        d, res = self._docx(lineless=True, ruled=True)
        self.assertGreaterEqual(len(d.tables), 4)       # раскладки + сетка + без линий + горизонтальные линии
        sec = d.sections[0]
        self.assertAlmostEqual(sec.page_width.pt, 595, delta=1)
        self.assertAlmostEqual(sec.left_margin.pt, 60, delta=2)

    def test_grid_table_cells(self):
        d, _ = self._docx()
        grid = max(d.tables, key=lambda t: len(t.rows))
        self.assertEqual((len(grid.rows), len(grid.columns)), (4, 5))
        self.assertEqual(grid.cell(1, 1).text, 'Laptop')
        self.assertEqual(grid.cell(3, 4).text, '8000')

    def test_paragraph_formatting_is_kept(self):
        d, _ = self._docx(grid=False)
        title = next(p for p in d.paragraphs if p.text == 'About the supplies')
        self.assertEqual(title.alignment, 1)             # center
        self.assertTrue(title.runs[0].bold)

    def test_merged_cells_roundtrip(self):
        html = ('<table class="gd-grid" style="width:200pt"><colgroup><col style="width:100pt"><col style="width:100pt"></colgroup>'
                '<tr><td colspan="2">Header</td></tr><tr><td>a</td><td>b</td></tr></table>')
        d = Document(io.BytesIO(html_to_docx(html)))
        t = d.tables[0]
        self.assertEqual(t.cell(0, 0).text, 'Header')
        self.assertEqual(t.cell(0, 0)._tc, t.cell(0, 1)._tc)

    def test_floating_stamp_is_anchored(self):
        doc = _letter(grid=False)
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
        pix.set_rect(pix.irect, (200, 30, 40))
        doc[0].insert_image(pymupdf.Rect(470, 680, 530, 740), stream=pix.tobytes('png'))   # печать поверх подписи
        res = _convert(doc)
        self.assertIn('data-float="1"', res.html)
        xml = Document(io.BytesIO(html_to_docx(res.html, res.meta)))._body._body.xml
        self.assertIn('<wp:anchor', xml)

    def test_empty_html_does_not_crash(self):
        Document(io.BytesIO(html_to_docx('')))


@skipUnless(HAS_TESSERACT, 'Tesseract не установлен')
class ScanTests(SimpleTestCase):
    def test_scanned_letter_keeps_tables_and_layout(self):
        res = _convert(_scan_of(_letter(lineless=True)))
        self.assertEqual(res.ocr_pages, 1)
        self.assertIn('class="gd-grid"', res.html)                      # таблица с линиями найдена на скане
        self.assertIn('Laptop', res.html)
        self.assertIn('class="gd-layout"', res.html)                    # реквизиты | адресат, подпись
        self.assertRegex(res.html, r'text-align:center[^"]*"><b>[^<]*supplies')   # заголовок по центру и жирный

    def test_sparse_page_is_recognised(self):
        """Короткое письмо (мало чернил) раньше теряло весь текст из-за автоконтраста."""
        doc = pymupdf.open()
        p = doc.new_page(width=595, height=842)
        _text(p, 60, 100, 'Short letter with very little ink on the page.')
        res = _convert(_scan_of(doc))
        self.assertIn('Short letter', res.html)


def _rotated(doc, angle):
    """Скан, повёрнутый целиком (как лист, поданный в сканер боком или вверх ногами)."""
    out = pymupdf.open()
    for pg in doc:
        pix = pg.get_pixmap(dpi=300, colorspace=pymupdf.csGRAY, alpha=False)
        im = Image.frombytes('L', (pix.width, pix.height), pix.samples).rotate(angle, expand=True)
        buf = io.BytesIO()
        im.save(buf, 'PNG')
        w, h = (pg.rect.height, pg.rect.width) if angle % 180 else (pg.rect.width, pg.rect.height)
        np_ = out.new_page(width=w, height=h)
        np_.insert_image(np_.rect, stream=buf.getvalue())
    return out


@skipUnless(HAS_TESSERACT and HAS_OCRMYPDF, 'Нужны Tesseract, OCRmyPDF и Ghostscript')
class OcrmypdfScanTests(SimpleTestCase):
    """Связка OCRmyPDF (поворот страниц) + Tesseract (слова, таблицы)."""

    @override_settings(OCRMYPDF_MODE='auto')
    def test_sideways_and_upside_down_scans_are_recognised(self):
        for angle in (90, 180):
            res = _convert(_rotated(_scan_of(_letter(lineless=True)), angle))
            self.assertTrue(res.ocrmypdf, angle)
            self.assertIn('Laptop', res.html, angle)
            self.assertIn('class="gd-grid"', res.html, angle)

    @override_settings(OCRMYPDF_MODE='off')
    def test_mode_off_does_not_call_ocrmypdf(self):
        with mock.patch.object(ocrprep, 'prepare') as prep:
            res = _convert(_scan_of(_letter(lineless=True)))
        prep.assert_not_called()
        self.assertFalse(res.ocrmypdf)


class OcrmypdfFallbackTests(SimpleTestCase):
    def _scan(self):
        doc = pymupdf.open()
        p = doc.new_page(width=595, height=842)
        _text(p, 60, 100, 'Short letter with very little ink on the page.')
        return _scan_of(doc)

    @skipUnless(HAS_TESSERACT, 'Tesseract не установлен')
    @override_settings(OCRMYPDF_MODE='auto')
    def test_auto_falls_back_to_builtin_prep_with_warning(self):
        with mock.patch.object(ocrprep, 'available', return_value=True), \
                mock.patch.object(ocrprep, 'prepare', side_effect=ocrprep.OcrPrepError('boom')):
            res = _convert(self._scan())
        self.assertFalse(res.ocrmypdf)
        self.assertIn('Short letter', res.html)
        self.assertTrue(any('OCRmyPDF не применён' in w for w in res.warnings))

    @skipUnless(HAS_TESSERACT, 'Tesseract не установлен')
    @override_settings(OCRMYPDF_MODE='auto')
    def test_auto_without_ocrmypdf_is_silent(self):
        with mock.patch.object(ocrprep, 'available', return_value=False):
            res = _convert(self._scan())
        self.assertFalse(res.ocrmypdf)
        self.assertEqual(res.warnings, [])

    @override_settings(OCRMYPDF_MODE='on')
    def test_mode_on_requires_ocrmypdf(self):
        with mock.patch.object(ocrprep, 'available', return_value=False):
            with self.assertRaises(ConversionError):
                _convert(self._scan())


class ViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('u', password='p')
        self.client.force_login(self.user)

    def test_convert_returns_page_meta_and_docx_accepts_it(self):
        pdf = io.BytesIO(_letter().tobytes())
        pdf.name = 'letter.pdf'
        r = self.client.post(reverse('convert'), {'pdf': pdf})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertIn('ml', data['page'])
        r2 = self.client.post(reverse('build_docx'), {'html': data['html'], 'filename': 'x', 'page': json.dumps(data['page'])})
        self.assertEqual(r2.status_code, 200)
        d = Document(io.BytesIO(r2.content))
        self.assertAlmostEqual(d.sections[0].page_width.pt, 595, delta=1)

    def test_docx_ignores_garbage_page_meta(self):
        r = self.client.post(reverse('build_docx'), {'html': '<p>hi</p>', 'filename': 'x', 'page': '{"w": "abc"}'})
        self.assertEqual(r.status_code, 200)


class LayoutUnitTests(SimpleTestCase):
    def test_xy_tree_splits_columns(self):
        def seg(x, y, t):
            return L.Seg([L.Word(x, y - 9, x + 6 * len(t), y + 3, t, size=12, base=y)], y)
        node = L.xy_tree([seg(60, 100, 'left one'), seg(300, 100, 'right one'), seg(60, 116, 'left two')])
        self.assertIsInstance(node, L.Cols)
        self.assertEqual(len(node.cols), 2)
