"""HTML редактора -> DOCX с сохранением вёрстки (отступы, выравнивание, интервалы, таблицы, колонки).

Понимает ровно тот HTML, который выдаёт letters/layout.py, плюс то, что получается при правке в редакторе
(div, br, b/i/u, ul/ol, h2/h3, font). Свой разбор вместо универсальных конвертеров: нужен точный контроль
над абзацем (exact-интервал, отступы в pt), таблицами (ширины колонок, границы, объединённые ячейки)
и «невидимыми» таблицами-раскладками (реквизиты слева / адресат справа, подпись и т. п.).
"""
import base64
import io
import re

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

DEFAULT_FONT = 'Times New Roman'
DEFAULT_SIZE = 12.0
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


# ----------------------------------------------------------------- разбор css
def parse_style(s):
    out = {}
    for part in (s or '').split(';'):
        if ':' in part:
            k, v = part.split(':', 1)
            out[k.strip().lower()] = v.strip()
    return out


def to_pt(v, base=DEFAULT_SIZE):
    """'12pt' / '16px' / '1.2em' / '12' -> pt (или None)."""
    if v is None:
        return None
    m = re.match(r'^\s*(-?\d+(?:\.\d+)?)\s*(pt|px|em|cm|mm|in|%)?\s*$', str(v))
    if not m:
        return None
    n, u = float(m.group(1)), m.group(2) or 'pt'
    return {'pt': n, 'px': n * .75, 'em': n * base, 'cm': n * 28.3465, 'mm': n * 2.83465,
            'in': n * 72, '%': n * base / 100}[u]


def font_name(css):
    if not css:
        return None
    first = css.split(',')[0].strip().strip('\'"')
    return first or None


# ------------------------------------------------------------ абзацы и runs --
class Ctx:
    def __init__(self, size=DEFAULT_SIZE, font=DEFAULT_FONT):
        self.size = size
        self.font = font


def _apply_run(run, fmt):
    f = run.font
    if fmt.get('b'):
        f.bold = True
    if fmt.get('i'):
        f.italic = True
    if fmt.get('u'):
        f.underline = True
    if fmt.get('size'):
        f.size = Pt(fmt['size'])
    name = fmt.get('font')
    if name:
        f.name = name
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is not None:
            for a in ('w:cs', 'w:eastAsia'):
                rf.set(qn(a), name)


def _runs(par, node, fmt, ctx):
    if isinstance(node, NavigableString):
        text = re.sub(r'\s+', ' ', str(node).replace('\xa0', ' '))
        if text:
            _apply_run(par.add_run(text), fmt)
        return
    if not isinstance(node, Tag):
        return
    n = node.name
    if n == 'br':
        par.add_run().add_break()
        return
    if n == 'img':
        _picture(par, node)
        return
    f = dict(fmt)
    if n in ('b', 'strong', 'th'):
        f['b'] = True
    elif n in ('i', 'em'):
        f['i'] = True
    elif n == 'u':
        f['u'] = True
    st = parse_style(node.get('style'))
    if st.get('font-weight') in ('bold', '700', '800', '900'):
        f['b'] = True
    if st.get('font-style') == 'italic':
        f['i'] = True
    if 'underline' in st.get('text-decoration', ''):
        f['u'] = True
    sz = to_pt(st.get('font-size'))
    if sz:
        f['size'] = sz
    fn = font_name(st.get('font-family'))
    if fn:
        f['font'] = fn
    if n == 'font' and node.get('size'):
        f['size'] = {'1': 8, '2': 10, '3': 12, '4': 14, '5': 18, '6': 24, '7': 36}.get(node.get('size'), f.get('size'))
    for c in node.children:
        _runs(par, c, f, ctx)


EMU_PT = 12700
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'


def _anchor_picture(run, stream, w_pt, h_pt, x_pt, dy_pt):
    """Плавающая картинка (печать/подпись поверх текста): X — от края страницы, Y — от верха абзаца."""
    shape = run.add_picture(stream, width=Pt(w_pt), height=Pt(h_pt))
    inline = shape._inline
    extent, doc_pr = inline.find(f'{{{WP}}}extent'), inline.find(f'{{{WP}}}docPr')
    cnv, graphic = inline.find(f'{{{WP}}}cNvGraphicFramePr'), inline[-1]
    anchor = OxmlElement('wp:anchor')
    for k, v in (('distT', '0'), ('distB', '0'), ('distL', '0'), ('distR', '0'), ('simplePos', '0'),
                 ('relativeHeight', '251658240'), ('behindDoc', '0'), ('locked', '0'), ('layoutInCell', '1'),
                 ('allowOverlap', '1')):
        anchor.set(k, v)
    sp = OxmlElement('wp:simplePos')
    sp.set('x', '0')
    sp.set('y', '0')
    anchor.append(sp)
    for tag, rel, off in (('wp:positionH', 'page', x_pt), ('wp:positionV', 'paragraph', dy_pt)):
        pos = OxmlElement(tag)
        pos.set('relativeFrom', rel)
        po = OxmlElement('wp:posOffset')
        po.text = str(int(round(off * EMU_PT)))
        pos.append(po)
        anchor.append(pos)
    anchor.append(extent)
    ee = OxmlElement('wp:effectExtent')
    for k in ('l', 't', 'r', 'b'):
        ee.set(k, '0')
    anchor.append(ee)
    anchor.append(OxmlElement('wp:wrapNone'))
    anchor.append(doc_pr)
    if cnv is not None:
        anchor.append(cnv)
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _picture(par, node):
    src = node.get('src', '')
    m = re.match(r'data:image/[\w.+-]+;base64,(.*)', src, re.S)
    if not m:
        return
    try:
        data = base64.b64decode(m.group(1))
        w = to_pt(node.get('width'), 1) or None
        h = to_pt(node.get('height'), 1) or None
        if node.get('data-float') and w and h and node.get('data-x') is not None:
            _anchor_picture(par.add_run(), io.BytesIO(data), w, h, float(node['data-x']), float(node.get('data-dy', 0)))
            return
        shape = par.add_run().add_picture(io.BytesIO(data), width=Pt(min(w, 520)) if w else None)
        for k in ('distT', 'distB', 'distL', 'distR'):        # без явных нулей LibreOffice добавляет поля ~9 pt
            shape._inline.set(k, '0')
    except Exception:  # noqa: BLE001  — битая картинка не должна ломать сохранение
        return


def _set_border_bottom(par, color='000000', sz=6, space=1):
    ppr = par._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single')
    bt.set(qn('w:sz'), str(sz))
    bt.set(qn('w:space'), str(int(max(1, min(31, round(space))))))
    bt.set(qn('w:color'), color)
    b.append(bt)
    ppr.append(b)


def fill_paragraph(par, node, ctx, default_after=None, base_style=None):
    st = dict(base_style or {})
    st.update(parse_style(node.get('style')))
    pf = par.paragraph_format
    base = to_pt(st.get('font-size')) or ctx.size
    font = font_name(st.get('font-family')) or ctx.font
    sb = to_pt(st.get('margin-top'), base)
    sa = to_pt(st.get('margin-bottom'), base)
    pf.space_before = Pt(sb or 0)
    pf.space_after = Pt(sa if sa is not None else (default_after if default_after is not None else 0))
    ml, mr, ti = to_pt(st.get('margin-left'), base), to_pt(st.get('margin-right'), base), to_pt(st.get('text-indent'), base)
    if ml:
        pf.left_indent = Pt(ml)
    if mr:
        pf.right_indent = Pt(mr)
    if ti:
        pf.first_line_indent = Pt(ti)
    al = st.get('text-align')
    if al in ALIGN:
        pf.alignment = ALIGN[al]
    lh = st.get('line-height')
    lhp = to_pt(lh, base) if lh and lh != 'normal' else None
    if lhp and re.search(r'(pt|px)\s*$', lh):
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(lhp)
    elif lhp:
        pf.line_spacing = float(lh) if re.match(r'^[\d.]+$', lh) else 1.0
    else:
        pf.line_spacing = 1.0
    if 'border-bottom' in st:
        _set_border_bottom(par, space=to_pt(st.get('padding-bottom'), base) or 1)
    fmt = {'size': base, 'font': font}
    if st.get('font-weight') in ('bold', '700'):
        fmt['b'] = True
    for c in node.children:
        _runs(par, c, fmt, Ctx(base, font))
    if not par.runs:
        # пустой абзац: размер шрифта задаёт высоту строки
        r = par.add_run('')
        r.font.size = Pt(base)


# --------------------------------------------------------------- таблицы -----
def _set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:shd'))
    if old is not None:
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color.lstrip('#').upper()[:6])
    tcPr.append(shd)


def _set_cell_borders(cell, kind, sides=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is not None:
        tcPr.remove(b)
    b = OxmlElement('w:tcBorders')
    if sides is None:
        sides = {'grid': ('top', 'left', 'bottom', 'right'), 'layout': ()}.get(kind, ('top', 'left', 'bottom', 'right'))
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        if side in sides:
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), '000000')
        else:
            e.set(qn('w:val'), 'nil')
        e.set(qn('w:space'), '0')
        b.append(e)
    tcPr.append(b)


def _cell_margins(table, top, left, bottom, right):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), str(int(v * 20)))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tblPr.append(m)


def _cell_pad(cell, top, left, bottom, right):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:tcMar'))
    if old is not None:
        tcPr.remove(old)
    m = OxmlElement('w:tcMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), str(int(v * 20)))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def _tbl_prop(table, ind_pt, width_pt):
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW')
        tblPr.append(w)
    w.set(qn('w:w'), str(int(width_pt * 20)))
    w.set(qn('w:type'), 'dxa')
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), str(int(ind_pt * 20)))
    ind.set(qn('w:type'), 'dxa')
    tblPr.append(ind)


def _int(v, default=1):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def _grid(table_node):
    """Раскладка ячеек: [(r, c, rowspan, colspan, td)], число строк, число колонок."""
    rows = [tr for tr in table_node.find_all('tr')]
    occupied, placed, ncols = set(), [], 0
    for r, tr in enumerate(rows):
        c = 0
        for td in tr.find_all(['td', 'th'], recursive=False):
            while (r, c) in occupied:
                c += 1
            rs = min(max(_int(td.get('rowspan')), 1), len(rows) - r)
            cs = min(max(_int(td.get('colspan')), 1), 60)
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
            placed.append((r, c, rs, cs, td))
            c += cs
            ncols = max(ncols, c)
    return placed, rows, ncols


def emit_table(container, node, ctx, add_table):
    cls = ' '.join(node.get('class', []))
    kind = 'layout' if 'gd-layout' in cls else 'custom' if 'gd-custom' in cls else 'grid'
    placed, rows, ncols = _grid(node)
    if not placed:
        return
    st = parse_style(node.get('style'))
    cols = [to_pt(c.get('style') and parse_style(c.get('style')).get('width')) for c in node.find_all('col')]
    if len(cols) != ncols or not all(cols):
        total = to_pt(st.get('width')) or 468.0
        cols = [total / ncols] * ncols
    t = add_table(len(rows), ncols)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    _tbl_prop(t, to_pt(st.get('margin-left')) or 0, sum(cols))
    _cell_margins(t, 0, 0, 0, 0)       # реальные отступы задаём в каждой ячейке (tcMar)
    for i, w in enumerate(cols):
        t.columns[i].width = Pt(w)
    for r, tr in enumerate(rows):
        h = to_pt(parse_style(tr.get('style')).get('height'))
        if h:
            t.rows[r].height = Pt(h)
            t.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST if kind != 'layout' else WD_ROW_HEIGHT_RULE.AT_LEAST
    for r, c, rs, cs, td in placed:
        cell = t.cell(r, c)
        if rs > 1 or cs > 1:
            cell = cell.merge(t.cell(r + rs - 1, c + cs - 1))
        cell.width = Pt(sum(cols[c:c + cs]))
        cst = parse_style(td.get('style'))
        if kind == 'custom':
            sides = tuple(sd for sd in ('top', 'right', 'bottom', 'left')
                          if cst.get(f'border-{sd}', 'none').strip() not in ('none', '0', ''))
            _set_cell_borders(cell, kind, sides)
        else:
            _set_cell_borders(cell, kind)
        if cst.get('white-space') == 'nowrap':
            nw = OxmlElement('w:noWrap')
            cell._tc.get_or_add_tcPr().append(nw)
        bg = cst.get('background') or cst.get('background-color')
        if bg and re.match(r'^#[0-9a-fA-F]{6}$', bg.strip()):
            _set_cell_shading(cell, bg.strip())
        if kind != 'layout' or any(k in cst for k in ('padding-left', 'padding-right', 'padding-top')):
            dl = 2.5 if kind != 'layout' else 0
            dt = 0.5 if kind != 'layout' else 0
            _cell_pad(cell, to_pt(cst.get('padding-top')) if 'padding-top' in cst else dt,
                      to_pt(cst.get('padding-left')) if 'padding-left' in cst else dl,
                      dt, to_pt(cst.get('padding-right')) if 'padding-right' in cst else dl)
        va = cst.get('vertical-align')
        if va:
            from docx.enum.table import WD_ALIGN_VERTICAL
            cell.vertical_alignment = {'middle': WD_ALIGN_VERTICAL.CENTER, 'bottom': WD_ALIGN_VERTICAL.BOTTOM,
                                       'top': WD_ALIGN_VERTICAL.TOP}.get(va, WD_ALIGN_VERTICAL.TOP)
        emit_cell(cell, td, ctx, cst, bold=td.name == 'th')
    return t


def emit_cell(cell, td, ctx, cst, bold=False):
    base = {k: v for k, v in cst.items() if k in ('font-size', 'font-family', 'text-align', 'line-height')}
    blocks = [c for c in td.children if isinstance(c, Tag) and c.name in ('p', 'div', 'table', 'ul', 'ol', 'h2', 'h3')]
    first = cell.paragraphs[0]
    if not blocks:
        wrap = BeautifulSoup('<p></p>', 'html.parser').p
        for ch in list(td.children):
            wrap.append(ch.extract() if hasattr(ch, 'extract') else ch)
        if bold:
            b = BeautifulSoup('<b></b>', 'html.parser').b
            for ch in list(wrap.children):
                b.append(ch.extract())
            wrap.append(b)
        fill_paragraph(first, wrap, ctx, default_after=0, base_style=base)
        return
    used_first = False
    for ch in td.children:
        if isinstance(ch, NavigableString):
            if str(ch).strip():
                p = first if not used_first else cell.add_paragraph()
                used_first = True
                w = BeautifulSoup('<p></p>', 'html.parser').p
                w.string = str(ch)
                fill_paragraph(p, w, ctx, default_after=0, base_style=base)
            continue
        if ch.name == 'table':
            emit_table(cell, ch, ctx, lambda r, c, cell=cell: cell.add_table(r, c))
            used_first = True
            continue
        p = first if not used_first else cell.add_paragraph()
        used_first = True
        fill_paragraph(p, ch, ctx, default_after=0, base_style=base)


# ----------------------------------------------------------------- документ ---
def html_to_docx(content_html: str, meta: dict | None = None) -> bytes:
    soup = BeautifulSoup(content_html or '', 'html.parser')
    for t in soup(['script', 'style', 'iframe', 'object']):
        t.decompose()
    for div in soup.find_all(['div', 'section']):
        if div.find('table'):
            div.unwrap()
    document = Document()
    sec = document.sections[0]
    m = meta or {}
    if m.get('w') and m.get('h'):
        sec.page_width, sec.page_height = Pt(m['w']), Pt(m['h'])
        sec.left_margin, sec.right_margin = Pt(m.get('ml', 56.7)), Pt(m.get('mr', 56.7))
        sec.top_margin, sec.bottom_margin = Pt(m.get('mt', 56.7)), Pt(m.get('mb', 28))
    else:
        sec.page_width, sec.page_height = Pt(595.3), Pt(841.9)  # A4
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(56.7)
    sec.header_distance = sec.footer_distance = Pt(0)
    normal = document.styles['Normal']
    normal.font.name = DEFAULT_FONT
    normal.font.size = Pt(DEFAULT_SIZE)
    npf = normal.paragraph_format
    npf.space_before = npf.space_after = Pt(0)
    npf.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is not None:
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), DEFAULT_FONT)

    ctx = Ctx()
    last = {'kind': None}

    def spacer(height_pt):
        """Пустой абзац заданной высоты — «отступ перед таблицей»."""
        p = document.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(max(height_pt, 1))
        r = p.add_run('')
        r.font.size = Pt(1)

    for node in list(soup.children):
        if isinstance(node, NavigableString):
            if str(node).strip():
                p = document.add_paragraph()
                w = BeautifulSoup('<p></p>', 'html.parser').p
                w.string = str(node)
                fill_paragraph(p, w, ctx, default_after=6)
                last['kind'] = 'p'
            continue
        n = node.name
        if n == 'table':
            st = parse_style(node.get('style'))
            sb = to_pt(st.get('margin-top')) or 0
            if sb >= 1 or last['kind'] == 'table':
                spacer(max(sb, 1))
            emit_table(document, node, ctx, lambda r, c: document.add_table(r, c))
            last['kind'] = 'table'
        elif n == 'hr':
            p = document.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(1)
            r = p.add_run()
            r.font.size = Pt(1)
            r.add_break(WD_BREAK.PAGE)
            last['kind'] = 'p'
        elif n in ('ul', 'ol'):
            for li in node.find_all('li', recursive=False):
                p = document.add_paragraph(style='List Bullet' if n == 'ul' else 'List Number')
                fill_paragraph(p, li, ctx, default_after=2)
            last['kind'] = 'p'
        elif n in ('h1', 'h2', 'h3', 'h4'):
            size = {'h1': 18, 'h2': 15, 'h3': 13, 'h4': 12}[n]
            p = document.add_paragraph()
            wrap = BeautifulSoup('<p></p>', 'html.parser').p
            b = BeautifulSoup('<b></b>', 'html.parser').b
            for ch in list(node.children):
                b.append(ch.extract())
            wrap.append(b)
            wrap['style'] = node.get('style', '')
            fill_paragraph(p, wrap, ctx, default_after=6, base_style={'font-size': f'{size}pt'})
            last['kind'] = 'p'
        elif n in ('p', 'div', 'section'):
            p = document.add_paragraph()
            fill_paragraph(p, node, ctx, default_after=6 if not parse_style(node.get('style')) else None)
            last['kind'] = 'p'
        elif n == 'img':
            p = document.add_paragraph()
            _picture(p, node)
            last['kind'] = 'p'
        elif n == 'br':
            continue
        else:
            p = document.add_paragraph()
            fill_paragraph(p, node, ctx, default_after=6)
            last['kind'] = 'p'

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
