"""Преобразование PDF (текстовых и сканов) в HTML для редактирования и в DOCX.

* Страницы с текстовым слоем разбираются напрямую через PyMuPDF
  (сохраняются абзацы, жирный/курсив, размеры шрифта, выравнивание, картинки).
* Страницы-сканы распознаются Tesseract OCR (rus+eng, 300 dpi, предобработка).
"""
import base64
import html
import io
import os
import statistics
from dataclasses import dataclass, field

import numpy as np
import pymupdf as fitz
from bs4 import BeautifulSoup
from django.conf import settings
from docx import Document
from docx.shared import Cm, Pt
from htmldocx import HtmlToDocx
from PIL import Image, ImageOps

MIN_TEXT_CHARS = 25  # меньше символов на странице -> считаем страницу сканом


class ConversionError(Exception):
    pass


@dataclass
class ConversionResult:
    html: str
    pages: int
    ocr_pages: int = 0
    warnings: list = field(default_factory=list)


def _setup_tesseract():
    import pytesseract

    if settings.TESSERACT_CMD:
        pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
    os.environ['TESSDATA_PREFIX'] = str(settings.TESSDATA_DIR)
    return pytesseract


def _align(x0, x1, page_w, left_margin, right_margin):
    width = x1 - x0
    if width > 0.8 * (right_margin - left_margin):
        return ''
    center = (x0 + x1) / 2
    if abs(center - page_w / 2) < page_w * 0.05 and x0 > left_margin + page_w * 0.05:
        return 'center'
    if x1 > right_margin - page_w * 0.03 and x0 > page_w * 0.45:
        return 'right'
    return ''


def _p(text_html, align='', tag='p', extra=''):
    style = f' style="text-align:{align}"' if align else ''
    return f'<{tag}{style}>{text_html}</{tag}>'


def _native_tables(page):
    """Таблицы текстового PDF: [(bbox, html)]."""
    out = []
    try:
        found = page.find_tables().tables
    except Exception:  # noqa: BLE001
        return out
    for t in found:
        rects = [c for c in t.cells if c]
        if len(rects) < 2 or t.col_count < 2:
            continue
        cells = [(*r, ' '.join(page.get_text('text', clip=pymupdf_rect(r)).split())) for r in rects]
        out.append((tuple(t.bbox), table_html(cells)))
    return out


def pymupdf_rect(r):
    return fitz.Rect(*r)


def _inside(bbox, outer, pad=2):
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return outer[0] - pad <= cx <= outer[2] + pad and outer[1] - pad <= cy <= outer[3] + pad


def _native_page(page):
    d = page.get_text('dict')
    tables = _native_tables(page)
    blocks = [b for b in d['blocks'] if b['type'] == 0 and not any(_inside(b['bbox'], t[0]) for t in tables)]
    page_w = page.rect.width
    xs0 = [b['bbox'][0] for b in blocks] or [0]
    xs1 = [b['bbox'][2] for b in blocks] or [page_w]
    lm, rm = min(xs0), max(xs1)
    sizes = [s['size'] for b in blocks for l in b['lines'] for s in l['spans'] if s['text'].strip()]
    base = statistics.median(sizes) if sizes else 11
    out = []  # (y_top, html)
    for tb, thtml in tables:
        out.append((tb[1], thtml))
    for b in d['blocks']:
        y_top = b['bbox'][1]
        if b['type'] == 1 and b.get('image'):
            ext = b.get('ext', 'png')
            w = min(int(b['bbox'][2] - b['bbox'][0]), 600)
            data = base64.b64encode(b['image']).decode()
            out.append((y_top, f'<p><img src="data:image/{ext};base64,{data}" width="{w}"></p>'))
            continue
        if b['type'] != 0 or any(_inside(b['bbox'], t[0]) for t in tables):
            continue
        parts, bsize, bold_all = [], [], True
        for l in b['lines']:
            line = []
            for s in l['spans']:
                t = s['text']
                if not t.strip():
                    line.append(html.escape(t))
                    continue
                bold = bool(s['flags'] & 16) or 'bold' in s['font'].lower()
                ital = bool(s['flags'] & 2) or 'italic' in s['font'].lower()
                bold_all &= bold
                bsize.append(s['size'])
                seg = html.escape(t)
                if bold:
                    seg = f'<b>{seg}</b>'
                if ital:
                    seg = f'<i>{seg}</i>'
                line.append(seg)
            parts.append(''.join(line).strip())
        body = ' '.join(p for p in parts if p).replace('</b> <b>', ' ').replace('</i> <i>', ' ')
        if not body:
            continue
        size = statistics.mean(bsize) if bsize else base
        align = _align(b['bbox'][0], b['bbox'][2], page_w, lm, rm)
        if size >= base * 1.35 and len(body) < 200:
            out.append((y_top, _p(body, align, 'h2')))
        elif size >= base * 1.15 and bold_all and len(body) < 200:
            out.append((y_top, _p(body, align, 'h3')))
        else:
            out.append((y_top, _p(body, align)))
    out.sort(key=lambda t: t[0])
    return [h for _, h in out]


# ------------------------------------------------------------------ таблицы
def _edges(values, tol=3.0):
    """Кластеризует координаты границ ячеек -> отсортированные уникальные значения."""
    out = []
    for v in sorted(values):
        if out and v - out[-1][-1] <= tol:
            out[-1].append(v)
        else:
            out.append([v])
    return [sum(g) / len(g) for g in out]


def _nearest(edges, v):
    return min(range(len(edges)), key=lambda i: abs(edges[i] - v))


def table_html(cells, tol=3.0):
    """cells: [(x0, y0, x1, y1, text)] -> <table> с colspan/rowspan по геометрии ячеек."""
    if not cells:
        return ''
    xs = _edges([c[0] for c in cells] + [c[2] for c in cells], tol)
    ys = _edges([c[1] for c in cells] + [c[3] for c in cells], tol)
    grid = {}
    for x0, y0, x1, y1, text in cells:
        c0, c1 = _nearest(xs, x0), _nearest(xs, x1)
        r0, r1 = _nearest(ys, y0), _nearest(ys, y1)
        if c1 <= c0 or r1 <= r0:
            continue
        grid[(r0, c0)] = (r1 - r0, c1 - c0, text)
    covered = set()
    rows_html = []
    for r in range(len(ys) - 1):
        tds = []
        for c in range(len(xs) - 1):
            if (r, c) in covered:
                continue
            if (r, c) in grid:
                rs, cs, text = grid[(r, c)]
                for dr in range(rs):
                    for dc in range(cs):
                        covered.add((r + dr, c + dc))
                attrs = (f' rowspan="{rs}"' if rs > 1 else '') + (f' colspan="{cs}"' if cs > 1 else '')
                tds.append(f'<td{attrs}>{html.escape(text)}</td>')
            else:
                tds.append('<td></td>')
        if tds:
            rows_html.append('<tr>' + ''.join(tds) + '</tr>')
    return '<table border="1" style="width:100%">' + ''.join(rows_html) + '</table>'


def _words_text(words):
    """words: [(x, y, w, h, text)] -> строки текста в порядке чтения."""
    words = sorted(words, key=lambda w: (w[1], w[0]))
    lines, cur, cur_y = [], [], None
    for w in words:
        if cur and abs(w[1] - cur_y) > 0.6 * max(w[3], 1):
            lines.append(cur)
            cur = []
        if not cur:
            cur_y = w[1]
        cur.append(w)
    if cur:
        lines.append(cur)
    return ' '.join(' '.join(x[4] for x in sorted(l)) for l in lines)


def _runs_mask(b, k, axis):
    """Пиксели, входящие в непрерывные отрезки True длиной >= k (axis=1 — по строкам)."""
    if axis == 0:
        return _runs_mask(b.T, k, 1).T
    h, w = b.shape
    if w < k:
        return np.zeros_like(b)
    c = np.zeros((h, w + 1), dtype=np.int32)
    np.cumsum(b, axis=1, out=c[:, 1:])
    start = (c[:, k:] - c[:, :-k]) == k          # окно из k подряд True, начиная с позиции i
    diff = np.zeros((h, w + 1), dtype=np.int8)
    diff[:, :w - k + 1] += start
    diff[:, k:] -= start
    return np.cumsum(diff, axis=1)[:, :w] > 0


def _dilate(b, r):
    out = b.copy()
    for d in range(1, r + 1):
        for ax in (0, 1):
            out |= np.roll(b, d, axis=ax)
            out |= np.roll(b, -d, axis=ax)
    return out


def _cluster_idx(idx):
    groups, cur = [], []
    for i in idx:
        if cur and i - cur[-1] > 2:
            groups.append(cur)
            cur = []
        cur.append(int(i))
    if cur:
        groups.append(cur)
    return [(g[0] + g[-1]) / 2 for g in groups]


def _components(mask_small):
    pts = {tuple(p) for p in np.argwhere(mask_small)}
    seen, comps = set(), []
    for p in pts:
        if p in seen:
            continue
        stack, ys, xs = [p], [], []
        seen.add(p)
        while stack:
            y, x = stack.pop()
            ys.append(y)
            xs.append(x)
            for dy in (-1, 0, 1):
                for dx in (-1, 0, 1):
                    q = (y + dy, x + dx)
                    if q in pts and q not in seen:
                        seen.add(q)
                        stack.append(q)
        comps.append((min(xs), min(ys), max(xs), max(ys)))
    return comps


def detect_scan_tables(arr, dpi):
    """Ищет таблицы с линиями на скане. arr — numpy uint8 (оттенки серого).

    Возвращает (tables, line_mask), tables = [{'bbox': (x0,y0,x1,y1), 'cells': [(x0,y0,x1,y1)]}].
    """
    dark = arr < 170
    k = max(int(dpi * 0.2), 20)
    hmask = _runs_mask(dark, k, 1)
    vmask = _runs_mask(dark, k, 0)
    lines = hmask | vmask
    f = 4
    H, W = lines.shape
    small = lines[:H // f * f, :W // f * f].reshape(H // f, f, W // f, f).any(axis=(1, 3))
    small = _dilate(small, 1)
    tables = []
    for sx0, sy0, sx1, sy1 in _components(small):
        x0, y0, x1, y1 = sx0 * f, sy0 * f, min((sx1 + 1) * f, W), min((sy1 + 1) * f, H)
        w, h = x1 - x0, y1 - y0
        if w < dpi * 1.0 or h < dpi * 0.3:
            continue
        hm, vm = hmask[y0:y1, x0:x1], vmask[y0:y1, x0:x1]
        ys = [y0 + v for v in _cluster_idx(np.where(hm.sum(axis=1) >= 0.3 * w)[0])]
        xs = [x0 + v for v in _cluster_idx(np.where(vm.sum(axis=0) >= 0.3 * h)[0])]
        if len(ys) < 2 or len(xs) < 2 or (len(ys) - 1) * (len(xs) - 1) < 2:
            continue

        def v_present(x, ya, yb):
            xi = int(round(x))
            seg = vmask[int(ya):int(yb), max(xi - 3, 0):xi + 4]
            return seg.size and seg.any(axis=1).mean() > 0.5

        def h_present(y, xa, xb):
            yi = int(round(y))
            seg = hmask[max(yi - 3, 0):yi + 4, int(xa):int(xb)]
            return seg.size and seg.any(axis=0).mean() > 0.5

        nr, nc = len(ys) - 1, len(xs) - 1
        done = [[False] * nc for _ in range(nr)]
        cells = []
        for i in range(nr):
            for j in range(nc):
                if done[i][j]:
                    continue
                j2 = j
                while j2 + 1 < nc and not v_present(xs[j2 + 1], ys[i] + 3, ys[i + 1] - 3):
                    j2 += 1
                i2 = i
                while i2 + 1 < nr and not h_present(ys[i2 + 1], xs[j] + 3, xs[j2 + 1] - 3):
                    i2 += 1
                for a in range(i, i2 + 1):
                    for b in range(j, j2 + 1):
                        done[a][b] = True
                cells.append((xs[j], ys[i], xs[j2 + 1], ys[i2 + 1]))
        tables.append({'bbox': (xs[0], ys[0], xs[-1], ys[-1]), 'cells': cells})
    return tables, lines


def _prepare_for_ocr(img):
    img = ImageOps.grayscale(img)
    img = ImageOps.autocontrast(img, cutoff=1)
    return img


def _ocr_page(page, pytesseract):
    dpi = settings.OCR_DPI
    pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY, alpha=False)
    img = Image.frombytes('L', (pix.width, pix.height), pix.samples)
    img = _prepare_for_ocr(img)
    arr = np.array(img)
    tables, line_mask = detect_scan_tables(arr, dpi)
    if tables:
        # стираем линии таблиц, чтобы они не мешали распознаванию текста
        clean = np.zeros_like(line_mask)
        for t in tables:
            x0, y0, x1, y1 = (int(v) for v in t['bbox'])
            clean[max(y0 - 4, 0):y1 + 5, max(x0 - 4, 0):x1 + 5] = True
        erase = _dilate(line_mask & clean, 2)
        arr = np.where(erase, 255, arr).astype(np.uint8)
        img = Image.fromarray(arr)
    try:
        data = pytesseract.image_to_data(
            img, lang=settings.OCR_LANGS, config='--oem 1 --psm 3',
            output_type=pytesseract.Output.DICT)
    except pytesseract.TesseractNotFoundError as e:
        raise ConversionError(
            'Не найден Tesseract OCR. Установите его (см. README.md) '
            'или задайте переменную TESSERACT_CMD.') from e
    except pytesseract.TesseractError as e:
        raise ConversionError(f'Ошибка распознавания: {e}') from e

    words = []  # (x, y, w, h, text, block, par, line)
    for i, txt in enumerate(data['text']):
        if not txt.strip() or float(data['conf'][i]) < 0:
            continue
        words.append((data['left'][i], data['top'][i], data['width'][i], data['height'][i], txt,
                      data['block_num'][i], data['par_num'][i], data['line_num'][i]))

    # слова внутри таблиц -> ячейки
    items = []  # (y_top, html)
    cell_words = {}
    rest = []
    for w in words:
        cx, cy = w[0] + w[2] / 2, w[1] + w[3] / 2
        placed = False
        for ti, t in enumerate(tables):
            tb = t['bbox']
            if tb[0] - 2 <= cx <= tb[2] + 2 and tb[1] - 2 <= cy <= tb[3] + 2:
                for ci, c in enumerate(t['cells']):
                    if c[0] <= cx <= c[2] and c[1] <= cy <= c[3]:
                        cell_words.setdefault((ti, ci), []).append(w[:5])
                        placed = True
                        break
            if placed:
                break
        if not placed:
            rest.append(w)
    for ti, t in enumerate(tables):
        cells = [(*c, _words_text(cell_words.get((ti, ci), []))) for ci, c in enumerate(t['cells'])]
        items.append((t['bbox'][1], table_html(cells, tol=dpi * 0.03)))

    # остальные слова: слово -> строка -> абзац
    paras = {}
    for x, y, w_, h_, txt, blk, par, ln in rest:
        paras.setdefault((blk, par), {}).setdefault(ln, []).append((x, y, w_, h_, txt))
    if paras:
        heights = [w[3] for p in paras.values() for l in p.values() for w in l]
        base = statistics.median(heights)
        page_w = pix.width
        lm = min(w[0] for p in paras.values() for l in p.values() for w in l)
        rm = max(w[0] + w[2] for p in paras.values() for l in p.values() for w in l)
        for key, lines in paras.items():
            pw = [w for l in lines.values() for w in l]
            text = ' '.join(' '.join(w[4] for w in sorted(l)) for _, l in sorted(lines.items()))
            x0, x1 = min(w[0] for w in pw), max(w[0] + w[2] for w in pw)
            h = statistics.median(w[3] for w in pw)
            align = _align(x0, x1, page_w, lm, rm)
            body = html.escape(text)
            y_top = min(w[1] for w in pw)
            if h >= base * 1.4 and len(text) < 200:
                items.append((y_top, _p(f'<b>{body}</b>', align, 'h2')))
            else:
                items.append((y_top, _p(body, align)))
    items.sort(key=lambda t: t[0])
    return [h for _, h in items]


def convert_pdf(data: bytes) -> ConversionResult:
    """data — содержимое PDF (без временных файлов: на Windows они мешают открытию)."""
    try:
        doc = fitz.open(stream=data, filetype='pdf')
    except Exception as e:  # noqa: BLE001
        raise ConversionError(f'Не удалось открыть PDF: {e}') from e
    if doc.needs_pass:
        raise ConversionError('PDF защищён паролем.')
    result = ConversionResult(html='', pages=doc.page_count)
    chunks = []
    pytesseract = None
    for n, page in enumerate(doc, start=1):
        native = _native_page(page) if len(page.get_text().strip()) >= MIN_TEXT_CHARS else []
        if native:
            items = native
        else:
            if pytesseract is None:
                pytesseract = _setup_tesseract()
            items = _ocr_page(page, pytesseract)
            result.ocr_pages += 1
            if not items:
                result.warnings.append(f'Страница {n}: текст не распознан.')
        if n > 1:
            chunks.append('<hr>')
        chunks.extend(items)
    doc.close()
    result.html = '\n'.join(chunks)
    return result


def _add_runs(paragraph, node, fmt=()):
    """Переносит inline-форматирование (b/i/u/br) из HTML-узла в абзац Word."""
    from bs4 import NavigableString
    if isinstance(node, NavigableString):
        text = str(node).replace('\n', ' ')
        if text:
            run = paragraph.add_run(text)
            run.bold = 'b' in fmt or None
            run.italic = 'i' in fmt or None
            run.underline = 'u' in fmt or None
        return
    if node.name == 'br':
        paragraph.add_run().add_break()
        return
    f = set(fmt)
    if node.name in ('b', 'strong', 'th'):
        f.add('b')
    elif node.name in ('i', 'em'):
        f.add('i')
    elif node.name == 'u':
        f.add('u')
    for child in node.children:
        _add_runs(paragraph, child, tuple(f))


def _add_table(document, table):
    """Таблица HTML -> таблица Word с границами, colspan и rowspan."""
    rows = table.find_all('tr')
    occupied, placed, n_cols = set(), [], 0
    for r, tr in enumerate(rows):
        c = 0
        for td in tr.find_all(['td', 'th'], recursive=False):
            while (r, c) in occupied:
                c += 1
            rs = max(int(td.get('rowspan', 1) or 1), 1)
            cs = max(int(td.get('colspan', 1) or 1), 1)
            rs = min(rs, len(rows) - r)
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
            placed.append((r, c, rs, cs, td))
            c += cs
            n_cols = max(n_cols, c)
    if not placed:
        return
    t = document.add_table(rows=len(rows), cols=n_cols)
    t.style = 'Table Grid'
    for r, c, rs, cs, td in placed:
        cell = t.cell(r, c)
        if rs > 1 or cs > 1:
            cell = cell.merge(t.cell(r + rs - 1, c + cs - 1))
        para = cell.paragraphs[0]
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        for child in td.children:
            _add_runs(para, child, ('b',) if td.name == 'th' else ())
    document.add_paragraph()


def html_to_docx(content_html: str) -> bytes:
    """Собирает DOCX из отредактированного пользователем HTML."""
    from docx.enum.text import WD_BREAK

    soup = BeautifulSoup(content_html, 'html.parser')
    for t in soup(['script', 'style', 'iframe', 'object']):
        t.decompose()
    for div in soup.find_all('div'):
        if not div.find('table'):
            div.name = 'p'
        else:
            div.unwrap()

    document = Document()
    sec = document.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    parser = HtmlToDocx()
    buffer = []

    def flush():
        chunk = ''.join(buffer).strip()
        buffer.clear()
        if chunk:
            parser.add_html_to_document(chunk, document)

    for node in list(soup.children):
        name = getattr(node, 'name', None)
        if name == 'table':
            flush()
            _add_table(document, node)
        elif name == 'hr':
            flush()
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            buffer.append(str(node))
    flush()

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
