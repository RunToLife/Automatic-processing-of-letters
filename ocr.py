"""Распознавание сканов PDF и построение WORD-документа."""
import io
import os

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from docx import Document
from docx.shared import Pt

import config

if config.TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_CMD


def pdf_to_images(pdf_path: str, dpi: int = None) -> list:
    """Рендерит каждую страницу PDF в изображение PIL с высоким DPI
    для максимально чёткого распознавания текста."""
    dpi = dpi or config.OCR_DPI
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    images = []
    doc = fitz.open(pdf_path)
    try:
        for page in doc:
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            images.append(img)
    finally:
        doc.close()
    return images


def ocr_image_to_text(image: Image.Image, languages: str = None) -> str:
    languages = languages or config.OCR_LANGUAGES
    # psm 6: считаем, что на странице единый блок текста - подходит для писем
    text = pytesseract.image_to_string(image, lang=languages, config="--psm 6")
    return text.strip()


def split_into_paragraphs(page_text: str) -> list:
    """Разбивает текст страницы на абзацы по пустым строкам."""
    raw_paragraphs = [p.strip() for p in page_text.split("\n\n")]
    paragraphs = []
    for p in raw_paragraphs:
        if not p:
            continue
        # склеиваем перенесённые строки внутри одного абзаца в одну строку
        lines = [line.strip() for line in p.split("\n") if line.strip()]
        paragraphs.append(" ".join(lines))
    return paragraphs


def recognize_pdf(pdf_path: str) -> list:
    """Возвращает список страниц, каждая - список абзацев (строк текста)."""
    images = pdf_to_images(pdf_path)
    pages = []
    for img in images:
        text = ocr_image_to_text(img)
        paragraphs = split_into_paragraphs(text)
        pages.append(paragraphs)
    return pages


def pages_to_editable_text(pages: list) -> str:
    """Склеивает распознанные страницы в единый редактируемый текст.
    Абзацы разделяются пустой строкой, страницы - меткой разрыва страницы."""
    page_texts = ["\n\n".join(page) for page in pages]
    return "\f".join(page_texts)


def editable_text_to_docx(text: str, out_path: str):
    """Строит DOCX-документ из отредактированного пользователем текста.
    Разрыв страницы кодируется символом \\f (form feed)."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    pages = text.split("\f")
    for i, page_text in enumerate(pages):
        if i > 0:
            document.add_page_break()
        paragraphs = [p for p in page_text.split("\n\n")]
        for para in paragraphs:
            para = para.strip("\n")
            if para == "":
                continue
            lines = para.split("\n")
            p = document.add_paragraph()
            for i, line in enumerate(lines):
                if i > 0:
                    p.add_run().add_break()
                p.add_run(line)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    document.save(out_path)


def render_page_previews(pdf_path: str, max_width: int = 900) -> list:
    """Возвращает список PNG-изображений страниц (в байтах) для показа в браузере."""
    dpi = 150
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    result = []
    doc = fitz.open(pdf_path)
    try:
        for page in doc:
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            buf = io.BytesIO(pix.tobytes("png"))
            result.append(buf.getvalue())
    finally:
        doc.close()
    return result
